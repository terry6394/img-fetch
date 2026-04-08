"""Tests for report_writer module."""

import tempfile
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from img_fetch.core.product import Product, FetchAttempt, ImageMetadata
from img_fetch.writers.report_writer import ReportWriter


@pytest.fixture
def temp_output_dir():
    """Create a temporary output directory."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture
def sample_products():
    """Create sample products for testing."""
    products = [
        Product(
            excel_row=1,
            name="Alpha Jacket",
            spec="Black-L",
            brand="HAGLOFS",
            status="success",
            image_path="/images/HAGLOFS/HAGLOFS_Alpha_Jacket_Black-L.jpg",
            tried_sources=["brand_site"],
            attempts=[
                FetchAttempt(source="brand_site", status="success", url="https://haglofs.com/product/123"),
            ],
            image_metadata=ImageMetadata(
                source_url="https://haglofs.com/images/alpha.jpg",
                source_page_title="Alpha Jacket - HAGLOFS",
                source_domain="haglofs.com",
                page_description="Professional outdoor jacket",
                page_keywords=["jacket", "outdoor", "haglofs"],
                alt_text=["Alpha Jacket front view", "Alpha Jacket back view"],
                image_width=1200,
                image_height=800,
                captured_at=datetime(2024, 1, 15, 10, 30, 0),
            ),
        ),
        Product(
            excel_row=2,
            name="Beta Vest",
            spec="Red-M",
            brand="HAGLOFS",
            status="failed",
            error_message="Image not found on any source",
            tried_sources=["brand_site", "ecommerce", "image_search"],
            attempts=[
                FetchAttempt(source="brand_site", status="failed", error="Not found"),
                FetchAttempt(source="ecommerce", status="failed", error="No image available"),
            ],
        ),
        Product(
            excel_row=3,
            name="Gamma Pants",
            spec="Blue-XL",
            brand="HAGLOFS",
            status="pending",
            tried_sources=[],
            attempts=[],
        ),
    ]
    return products


class TestReportWriter:
    """Tests for ReportWriter class."""

    def test_write_report_creates_file(self, temp_output_dir, sample_products):
        """Test that write_report creates an Excel file."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(sample_products)

        report_path = temp_output_dir / "report.xlsx"
        assert report_path.exists()

    def test_report_has_correct_sheets(self, temp_output_dir, sample_products):
        """Test that report has required sheets."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(sample_products)

        wb = load_workbook(temp_output_dir / "report.xlsx")
        sheet_names = wb.sheetnames

        assert "Product Images" in sheet_names
        assert "Statistics" in sheet_names

    def test_report_header_row(self, temp_output_dir, sample_products):
        """Test that header row contains all required columns."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(sample_products)

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Product Images"]

        headers = [cell.value for cell in ws[1]]
        expected_headers = [
            "Excel Row", "商品名称", "商品规格", "品牌", "Status",
            "Image Path", "Source Domain", "Source URL", "Source Page Title",
            "Page Description", "Page Keywords", "Alt Text", "Captured At",
            "Error Message", "Tried Sources",
        ]

        for expected in expected_headers:
            assert expected in headers

    def test_report_data_rows(self, temp_output_dir, sample_products):
        """Test that data rows are written correctly."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(sample_products)

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Product Images"]

        # Row 2 - first product
        row2 = [cell.value for cell in ws[2]]
        assert row2[0] == 1  # excel_row
        assert row2[1] == "Alpha Jacket"  # name
        assert row2[2] == "Black-L"  # spec
        assert row2[3] == "HAGLOFS"  # brand
        assert row2[4] == "success"  # status

        # Row 3 - second product (failed)
        row3 = [cell.value for cell in ws[3]]
        assert row3[0] == 2
        assert row3[1] == "Beta Vest"
        assert row3[4] == "failed"
        assert row3[13] == "Image not found on any source"  # error_message
        assert row3[14] == "brand_site, ecommerce, image_search"  # tried_sources

    def test_report_includes_metadata(self, temp_output_dir, sample_products):
        """Test that successful products include image metadata."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(sample_products)

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Product Images"]

        # Row 2 - product with metadata
        row2 = [cell.value for cell in ws[2]]
        assert row2[6] == "haglofs.com"  # source_domain
        assert row2[7] == "https://haglofs.com/images/alpha.jpg"  # source_url
        assert row2[8] == "Alpha Jacket - HAGLOFS"  # source_page_title
        assert row2[9] == "Professional outdoor jacket"  # page_description
        assert "jacket" in row2[10]  # page_keywords
        assert "Alpha Jacket front view" in row2[11]  # alt_text
        assert row2[12] == "2024-01-15 10:30:00"  # captured_at

    def test_report_pending_product_has_empty_metadata(self, temp_output_dir, sample_products):
        """Test that pending products have empty metadata fields."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(sample_products)

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Product Images"]

        # Row 4 - pending product
        row4 = [cell.value for cell in ws[4]]
        assert row4[0] == 3
        assert row4[4] == "pending"
        # Metadata fields should be empty strings
        assert row4[6] == ""  # source_domain
        assert row4[7] == ""  # source_url
        assert row4[8] == ""  # source_page_title

    def test_statistics_sheet_content(self, temp_output_dir, sample_products):
        """Test that statistics sheet contains correct data."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(sample_products)

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Statistics"]

        # Check title
        assert ws["A1"].value == "Image Fetch Task Statistics"

        # Check statistics values (starting at row 5)
        stats = {}
        for row in range(5, 10):
            metric = ws.cell(row=row, column=1).value
            value = ws.cell(row=row, column=2).value
            if metric:
                stats[metric] = value

        assert stats["Total Products"] == 3
        assert stats["Success"] == 1
        assert stats["Failed"] == 1
        assert stats["Pending"] == 1
        assert stats["Skipped"] == 0

    def test_statistics_sheet_has_source_stats(self, temp_output_dir, sample_products):
        """Test that statistics sheet includes source usage statistics."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(sample_products)

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Statistics"]

        # Find source statistics section
        source_section_row = None
        for row in range(1, 20):
            if ws.cell(row=row, column=1).value == "Source Usage Statistics":
                source_section_row = row
                break

        assert source_section_row is not None

        # Check that sources are listed
        source_col = []
        for row in range(source_section_row + 1, source_section_row + 10):
            source = ws.cell(row=row, column=1).value
            if source and source != "Source":
                source_col.append(source)

        assert "brand_site" in source_col
        assert "ecommerce" in source_col

    def test_report_path(self, temp_output_dir, sample_products):
        """Test that get_report_path returns correct path."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(sample_products)

        report_path = writer.get_report_path()
        assert report_path == temp_output_dir / "report.xlsx"


class TestReportWriterEdgeCases:
    """Edge case tests for ReportWriter."""

    def test_write_report_with_empty_products(self, temp_output_dir):
        """Test writing report with no products."""
        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report([])

        report_path = temp_output_dir / "report.xlsx"
        assert report_path.exists()

        wb = load_workbook(report_path)
        ws = wb["Product Images"]

        # Should have header row only
        assert ws.max_row == 1

    def test_write_report_with_product_no_metadata(self, temp_output_dir):
        """Test writing report for product without image metadata."""
        product = Product(
            excel_row=1,
            name="Test Product",
            spec="Spec1",
            brand="TEST",
            status="success",
        )

        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report([product])

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Product Images"]

        row = [cell.value for cell in ws[2]]
        assert row[0] == 1
        assert row[1] == "Test Product"
        # Metadata fields should be empty
        assert row[6] == ""
        assert row[7] == ""
        assert row[12] == ""  # captured_at

    def test_write_report_with_all_statuses(self, temp_output_dir):
        """Test report with products having all different statuses."""
        products = [
            Product(excel_row=1, name="P1", spec="", brand="B", status="success"),
            Product(excel_row=2, name="P2", spec="", brand="B", status="failed"),
            Product(excel_row=3, name="P3", spec="", brand="B", status="pending"),
            Product(excel_row=4, name="P4", spec="", brand="B", status="skipped"),
        ]

        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(products)

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Product Images"]

        statuses = [ws.cell(row=r, column=5).value for r in range(2, 6)]
        assert statuses == ["success", "failed", "pending", "skipped"]

    def test_keywords_and_alt_as_comma_separated(self, temp_output_dir):
        """Test that keywords and alt text are stored as comma-separated strings."""
        product = Product(
            excel_row=1,
            name="Test",
            spec="",
            brand="B",
            status="success",
            image_metadata=ImageMetadata(
                source_url="http://example.com/img.jpg",
                page_keywords=["keyword1", "keyword2", "keyword3"],
                alt_text=["alt1", "alt2"],
            ),
        )

        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report([product])

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Product Images"]

        row = [cell.value for cell in ws[2]]
        assert row[10] == "keyword1, keyword2, keyword3"  # page_keywords
        assert row[11] == "alt1, alt2"  # alt_text

    def test_multiple_brands_in_statistics(self, temp_output_dir):
        """Test statistics when products have different brands."""
        products = [
            Product(excel_row=1, name="P1", spec="", brand="HAGLOFS", status="success", tried_sources=["brand_site"]),
            Product(excel_row=2, name="P2", spec="", brand="HAGLOFS", status="success", tried_sources=["brand_site"]),
            Product(excel_row=3, name="P3", spec="", brand="CANADA_GOOSE", status="failed", tried_sources=["ecommerce"]),
            Product(excel_row=4, name="P4", spec="", brand="STONE_ISLAND", status="pending", tried_sources=[]),
        ]

        writer = ReportWriter(output_dir=temp_output_dir)
        writer.write_report(products)

        wb = load_workbook(temp_output_dir / "report.xlsx")
        ws = wb["Statistics"]

        # Find source statistics
        stats = {}
        for row in range(14, 20):
            source = ws.cell(row=row, column=1).value
            count = ws.cell(row=row, column=2).value
            if source and count is not None:
                stats[source] = count

        # brand_site should have 2 (from 2 HAGLOFS products)
        assert stats.get("brand_site") == 2
        # ecommerce should have 1 (from 1 CANADA_GOOSE product)
        assert stats.get("ecommerce") == 1
